#!/usr/bin/env python3
"""
Apply confirmed subject->(Category, Type) mappings to the Classification sheet of
an EIP workbook. For each data row whose subject matches a rule, set category,
type, quantity (by convention), and mark the row confirmed. Rows that match no
rule are left as-is (typically needs-review) so nothing is silently guessed.

This is a convenience for bulk-applying a user's decisions; the user can always
edit the sheet by hand instead.

Usage:
    python3 apply_mappings.py <workbook.xlsx> <mappings.json>

mappings.json: a list of
    {"pattern": "<regex over subject>", "category": "...", "type": "...",
     "quantity_mode": "hours"|"count"}
Order matters: the first matching rule wins.
"""
import json
import re
import sys

import openpyxl
from openpyxl.styles import PatternFill


def main():
    if len(sys.argv) != 3:
        sys.exit("usage: apply_mappings.py <workbook.xlsx> <mappings.json>")
    wb_path, map_path = sys.argv[1], sys.argv[2]
    rules_raw = json.load(open(map_path))
    rules = [(re.compile(r["pattern"], re.I), r["category"], r["type"],
              r.get("quantity_mode", "count")) for r in rules_raw]

    wb = openpyxl.load_workbook(wb_path)
    sheet = wb["Classification"]
    header = [c.value for c in sheet[1]]
    col = {h: i for i, h in enumerate(header)}
    clear = PatternFill(fill_type=None)

    applied = 0
    for row in sheet.iter_rows(min_row=2):
        subject = row[col["subject"]].value or ""
        duration = row[col["duration_hours"]].value
        for rx, category, type_name, qmode in rules:
            if rx.search(subject):
                row[col["category"]].value = category
                row[col["type"]].value = type_name
                if qmode == "hours":
                    row[col["quantity"]].value = max(1, round(float(duration or 0)))
                else:
                    row[col["quantity"]].value = 1
                row[col["status"]].value = "confirmed"
                row[col["needs_review"]].value = "no"
                row[col["confidence"]].value = "user-confirmed"
                for cell in row:
                    cell.fill = clear
                applied += 1
                break
    wb.save(wb_path)
    print(f"Applied mappings to {applied} rows in {wb_path}.")


if __name__ == "__main__":
    main()
