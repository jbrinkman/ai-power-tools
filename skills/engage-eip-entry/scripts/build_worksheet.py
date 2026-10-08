#!/usr/bin/env python3
"""
Phase 1 helper: read an EIP calendar spreadsheet and add a `Classification`
worksheet INTO THE SAME .xlsx file, with a best-guess Category/Type/Quantity, a
confidence level, and a needs_review flag for every data row. No rows are dropped
— the input sheet is already the eligible set. The user edits the Classification
sheet (fixing mappings, flipping status to `confirmed`) before Phase 2 enters the
values into Engage. Keeping the classification in the same workbook keeps all the
data together in one file.

Usage:
    python3 build_worksheet.py <workbook.xlsx> \
        [--source-sheet "EIP Entries"] [--class-sheet "Classification"] \
        [--rules rules.json]

By default this edits the workbook IN PLACE, adding (or replacing) the
Classification sheet. The source data sheet is left untouched.

quantity_mode:
    "hours"  -> Quantity = round(duration_hours) (min 1); raw duration kept for review
    "count"  -> Quantity = 1 (one entry per occurrence)
    "manual" -> left blank for the user to fill
"""
import argparse
import json
import os
import re
import sys

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

# Conservative keyword rules. A rule fires only when the subject clearly implies a
# single catalog Type. Ambiguous rows — especially internal recurring meetings —
# are left for review rather than guessed.
# (pattern, category, type, quantity_mode, confidence, note)
RULES = [
    (r"givecamp|improving cares|community service|habitat|volunteer",
     "Improving Cares", "Community Service", "count", "high",
     "Community service / Improving Cares event."),
    (r"\bAIDV\b|ai for developers|core ai skills|spec[- ]driven development|ai workshop|improvingu .*deliver",
     "Education/Coaching", "ImprovingU Instructor Delivery", "hours", "medium",
     "Assumed you INSTRUCTED an ImprovingU session (quantity = duration hours). "
     "If you merely attended, use 'ImprovingU Attendance'; if external/client, reclassify."),
    (r"town ?hall|come together|movie night|fantasy draft|ice cream|congrats",
     "Come Together", "Virtual Attendance", "count", "medium",
     "Assumed Come Together / Virtual Attendance. If attended IN PERSON, switch Type "
     "to 'In-Person Attendance' (different point value)."),
    (r"\b(spoke|speaking|presented|presenting|presentation|talk|lightning talk)\b",
     "Industry Contribution/Leadership", "Presentation - User Group", "count", "medium",
     "Assumed you PRESENTED (speaker). Confirm the exact Presentation type "
     "(User Group / Conference / Major / Improving Talks) and fill Topic Area + Attendance."),
    (r"\b(party|celebration|holiday party|happy hour|social|mixer|improving event)\b",
     "Networking", "Improving Event Attendance", "count", "medium",
     "Assumed attendance at an Improving event. Set the Attendance Type (Virtual / In Person)."),
    (r"user group|meetup|roundtable|community meeting|valkey",
     "Industry Participation", "User/Professional Group Attendance", "count", "medium",
     "Assumed external group ATTENDANCE. If you PRESENTED, reclassify under Industry "
     "Contribution/Leadership; if you LEAD/CONTRIBUTE (open source), use that type."),
]
COMPILED = [(re.compile(p, re.I), c, t, q, conf, note) for (p, c, t, q, conf, note) in RULES]

HEADERS = ["status", "reporting_period", "date", "subject", "duration_hours",
           "category", "type", "quantity", "notes", "org_url", "extra_fields_required",
           "confidence", "needs_review", "reason"]

# Required extra fields per (category, type), loaded from the bundled matrix so the
# worksheet tells Phase 2 what else to fill. Keyed "Category ||| Type".
_MATRIX_PATH = os.path.join(os.path.dirname(__file__), "..", "references",
                            "extra-fields-matrix.json")


def load_extra_matrix():
    try:
        raw = json.load(open(_MATRIX_PATH))
        data = raw.get("result", raw)
        out = {}
        for combo, fields in data.items():
            cat, typ = [p.strip() for p in combo.split("|||")]
            out[(cat, typ)] = [f["label"] for f in fields if f.get("required")]
        return out
    except Exception:
        return {}


EXTRA_MATRIX = load_extra_matrix()


def extra_fields_for(category, type_name):
    return ", ".join(EXTRA_MATRIX.get((category, type_name), []))


def quarter_of(dt):
    return f"{dt.year}-Q{(dt.month - 1) // 3 + 1}"


def classify(subject):
    for rx, cat, typ, qmode, conf, note in COMPILED:
        if rx.search(subject):
            return cat, typ, qmode, conf, note
    return "", "", "manual", "none", (
        "No confident match. Likely an internal recurring meeting (huddle, 1:1, "
        "sync, exec meeting) whose EIP category/type you must decide. Fill Category "
        "+ Type (see references/catalog.json) or set status appropriately.")


def quantity_for(mode, duration_hours):
    if mode == "hours":
        return max(1, round(float(duration_hours or 0)))
    if mode == "count":
        return 1
    return ""


def find_header_row(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = [str(c).strip().lower() if c is not None else "" for c in row]
        if "subject" in vals and any(v.startswith("start") for v in vals):
            cols = {v: j for j, v in enumerate(vals)}
            return i, cols
    return None, None


def col_index(cols, *prefixes):
    for p in prefixes:
        for key, idx in cols.items():
            if key.startswith(p):
                return idx
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--source-sheet", default="EIP Entries")
    ap.add_argument("--class-sheet", default="Classification")
    ap.add_argument("--rules")
    ap.add_argument("--out", help="write to a different file instead of in place")
    args = ap.parse_args()

    if args.rules and os.path.exists(args.rules):
        ext = json.load(open(args.rules))
        COMPILED[:] = [(re.compile(r["pattern"], re.I), r["category"], r["type"],
                        r["quantity_mode"], r.get("confidence", "medium"),
                        r.get("note", "")) for r in ext]

    wb = openpyxl.load_workbook(args.workbook)  # keep formulas/formatting of other sheets
    src = wb[args.source_sheet] if args.source_sheet in wb.sheetnames else wb.worksheets[0]

    # data_only read for values (a second load) so we get computed dates/durations
    wb_vals = openpyxl.load_workbook(args.workbook, data_only=True)
    src_vals = wb_vals[src.title]

    header_row, cols = find_header_row(src_vals)
    if header_row is None:
        sys.exit("Could not find a header row with 'Subject' and 'Start' in sheet "
                 f"'{src.title}'.")
    ci_subj = col_index(cols, "subject")
    ci_start = col_index(cols, "start", "date")
    ci_dur = col_index(cols, "duration")

    out_rows = []
    for row in src_vals.iter_rows(min_row=header_row + 1, values_only=True):
        if ci_subj >= len(row) or row[ci_subj] is None:
            continue
        subject = str(row[ci_subj]).strip()
        start = row[ci_start]
        duration = row[ci_dur] if ci_dur is not None and ci_dur < len(row) else None
        date_str = start.strftime("%m/%d/%Y") if hasattr(start, "strftime") else str(start)
        period = quarter_of(start) if hasattr(start, "year") else ""
        cat, typ, qmode, conf, note = classify(subject)
        qty = quantity_for(qmode, duration)
        extra = extra_fields_for(cat, typ)
        # A row that needs extra fields but has no value for them yet should be
        # reviewed, so Phase 2 is never asked to submit an incomplete form.
        needs_extra = bool(extra)
        review = conf in ("none", "medium") or needs_extra
        reason = note
        if needs_extra:
            reason = (note + " " if note else "") + (
                f"Requires extra field(s): {extra} — supply the value(s) "
                "(e.g. org_url) before entry.")
        out_rows.append({
            "status": "needs-review" if review else "confirmed",
            "reporting_period": period,
            "date": date_str,
            "subject": subject,
            "duration_hours": round(float(duration), 2) if duration else "",
            "category": cat,
            "type": typ,
            "quantity": qty,
            "notes": subject,
            "org_url": "",
            "extra_fields_required": extra,
            "confidence": conf,
            "needs_review": "yes" if review else "no",
            "reason": reason,
        })

    # Replace an existing Classification sheet so re-runs are idempotent.
    if args.class_sheet in wb.sheetnames:
        wb.remove(wb[args.class_sheet])
    cls = wb.create_sheet(args.class_sheet)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2d2471")
    review_fill = PatternFill("solid", fgColor="FFF3CD")  # pale amber for needs-review
    cls.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = cls.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")

    for r in out_rows:
        cls.append([r[h] for h in HEADERS])
        if r["needs_review"] == "yes":
            excel_row = cls.max_row
            for c in range(1, len(HEADERS) + 1):
                cls.cell(row=excel_row, column=c).fill = review_fill

    # Column widths for readability.
    widths = {"status": 13, "reporting_period": 16, "date": 12, "subject": 42,
              "duration_hours": 14, "category": 28, "type": 34, "quantity": 10,
              "notes": 42, "org_url": 32, "extra_fields_required": 24, "confidence": 12, "needs_review": 13, "reason": 70}
    for i, h in enumerate(HEADERS, start=1):
        cls.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths.get(h, 16)
    cls.freeze_panes = "A2"

    # A status dropdown makes the review step less error-prone.
    dv = DataValidation(type="list", formula1='"needs-review,confirmed,skip"', allow_blank=True)
    cls.add_data_validation(dv)
    dv.add(f"A2:A{cls.max_row}")

    out_path = args.out or args.workbook
    wb.save(out_path)

    total = len(out_rows)
    review = sum(1 for r in out_rows if r["needs_review"] == "yes")
    print(f"Added '{args.class_sheet}' sheet to {out_path}: {total} rows "
          f"({review} need review, {total - review} high-confidence). "
          f"Source sheet '{src.title}' left untouched.")


if __name__ == "__main__":
    main()
